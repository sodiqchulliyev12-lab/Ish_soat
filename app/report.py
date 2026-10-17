"""Attendance math, text summaries and the monthly Excel file."""
import io
from dataclasses import dataclass, field
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import db
from .config import COMPANY, LEAVE_KINDS, PAID_LEAVE
from .days import Calendar, schedule
from .utils import MONTHS, WEEKDAYS, fmt_minutes, hm, minutes, money, now, parse_stamp, today


@dataclass
class ShiftInfo:
    day: date
    in_at: str
    out_at: str | None
    worked: int          # minutes
    late: int            # minutes
    early: int           # minutes
    auto: bool
    note: str
    branch: str


def branch_names() -> dict[int, str]:
    return {b["id"]: b["name"] for b in db.branches()}


def shift_info(s, settings: dict, worker=None, branches: dict | None = None) -> ShiftInfo:
    worker = worker if worker is not None else db.get_worker(s["tg_id"])
    day = date.fromisoformat(s["date"])
    start, end = schedule(worker, settings).window(day)
    t_in = parse_stamp(s["in_at"])
    late = minutes(start, t_in) if minutes(start, t_in) > int(settings["grace"]) else 0
    if s["out_at"]:
        t_out = parse_stamp(s["out_at"])
        worked = max(minutes(t_in, t_out), 0)
        early = max(minutes(t_out, end), 0) if not s["auto_closed"] else 0
    else:
        worked = max(minutes(t_in, now()), 0)
        early = 0
    branch = (branches if branches is not None else branch_names()).get(s["branch_id"], "")
    return ShiftInfo(day, s["in_at"], s["out_at"], worked, late, early, bool(s["auto_closed"]), s["note"] or "", branch)


@dataclass
class WorkerMonth:
    worker: object
    sched_label: str = ""
    branch: str = ""
    workdays: int = 0            # workdays so far in the period (after the worker joined)
    month_workdays: int = 0      # all workdays of the month, for prorating monthly salary
    shifts: list[ShiftInfo] = field(default_factory=list)
    absent_days: list[date] = field(default_factory=list)
    leave_days: dict[str, list[date]] = field(default_factory=dict)

    @property
    def days(self) -> int:
        return len(self.shifts)

    @property
    def worked(self) -> int:
        return sum(s.worked for s in self.shifts)

    @property
    def late_count(self) -> int:
        return sum(1 for s in self.shifts if s.late)

    @property
    def late_min(self) -> int:
        return sum(s.late for s in self.shifts)

    @property
    def early_min(self) -> int:
        return sum(s.early for s in self.shifts)

    @property
    def paid_leave(self) -> int:
        return sum(len(v) for k, v in self.leave_days.items() if k in PAID_LEAVE)

    @property
    def unpaid_leave(self) -> int:
        return sum(len(v) for k, v in self.leave_days.items() if k not in PAID_LEAVE)

    @property
    def pay(self) -> float | None:
        w = self.worker
        if not w["rate"] or not w["rate_type"]:
            return None
        if w["rate_type"] == "oylik":
            paid_days = min(self.days + self.paid_leave, self.month_workdays)
            return w["rate"] * paid_days / max(self.month_workdays, 1)
        if w["rate_type"] == "kunlik":
            return w["rate"] * (self.days + self.paid_leave)
        return w["rate"] * self.worked / 60


def month_stats(start: date, end: date, only: int | None = None) -> list[WorkerMonth]:
    settings = db.get_settings()
    cal = Calendar(start, end, settings)
    names = branch_names()
    people = [w for w in db.workers("active") if only is None or w["tg_id"] == only]
    shifts = db.shifts_between(start, end, only)
    current = now()
    result = []
    for w in people:
        sched = schedule(w, settings)
        wm = WorkerMonth(worker=w, sched_label=sched.label(), branch=names.get(w["branch_id"], "Hammasi" if names else ""))
        joined = date.fromisoformat(w["created_at"][:10])
        mine = {s["date"]: s for s in shifts if s["tg_id"] == w["tg_id"]}
        wm.shifts = [shift_info(s, settings, w, names) for s in mine.values()]
        day = start
        while day < end:
            kind = cal.kind(w, day)
            if kind != "holiday" and kind != "weekend":
                wm.month_workdays += 1
            finished = current >= sched.window(day)[1]
            if day >= joined and finished:
                if kind.startswith("leave:"):
                    wm.leave_days.setdefault(kind[6:], []).append(day)
                elif kind == "work":
                    wm.workdays += 1
                    if day.isoformat() not in mine:
                        wm.absent_days.append(day)
            day += timedelta(days=1)
        result.append(wm)
    return result


# ---------- texts ----------

def month_title(start: date) -> str:
    return f"{MONTHS[start.month]} {start.year}"


def worker_month_text(wm: WorkerMonth, start: date) -> str:
    lines = [f"📊 <b>{wm.worker['name']}: {month_title(start)}</b>",
             f"🕘 Grafik: {wm.sched_label}",
             f"✅ Kelgan kunlar: {wm.days}",
             f"❌ Kelmagan kunlar: {len(wm.absent_days)}"]
    for kind, days in wm.leave_days.items():
        lines.append(f"🏖 {LEAVE_KINDS[kind]}: {len(days)} kun")
    lines += [f"⏱ Jami ishlangan: {fmt_minutes(wm.worked)}",
              f"⏰ Kechikishlar: {wm.late_count} marta ({fmt_minutes(wm.late_min)})"]
    if wm.early_min:
        lines.append(f"🏃 Erta ketish: {fmt_minutes(wm.early_min)}")
    if wm.pay is not None:
        lines.append(f"💰 Hisoblangan maosh: <b>{money(wm.pay)}</b>")
    if wm.absent_days:
        lines.append("Kelmagan kunlar: " + ", ".join(d.strftime("%d.%m") for d in wm.absent_days[-10:]))
    return "\n".join(lines)


def month_summary_text(stats: list[WorkerMonth], start: date) -> str:
    if not stats:
        return "Hali tasdiqlangan ishchi yo'q."
    lines = [f"📊 <b>Oylik hisobot: {month_title(start)}</b>", ""]
    total_pay = 0.0
    for wm in sorted(stats, key=lambda x: x.worker["name"]):
        pay = f" · {money(wm.pay)}" if wm.pay is not None else ""
        total_pay += wm.pay or 0
        leave = f", ta'til/ruxsat {wm.paid_leave + wm.unpaid_leave}" if wm.leave_days else ""
        lines.append(f"<b>{wm.worker['name']}</b>: {wm.days} kun, {fmt_minutes(wm.worked)}, "
                     f"kelmadi {len(wm.absent_days)}{leave}, kechikdi {wm.late_count}{pay}")
    holidays = db.holidays_between(start, (start + timedelta(days=32)).replace(day=1))
    if holidays:
        lines += ["", "🎉 Bayramlar: " + ", ".join(f"{date.fromisoformat(d):%d.%m} {n}" for d, n in holidays.items())]
    if total_pay:
        lines += ["", f"💰 Jami maosh: <b>{money(total_pay)}</b>"]
    lines += ["", "Batafsil Excel faylda 👇"]
    return "\n".join(lines)


def today_text() -> str:
    settings = db.get_settings()
    day, current = today(), now()
    cal = Calendar(day - timedelta(days=1), day + timedelta(days=1), settings)
    names = branch_names()
    people = db.workers("active")
    if not people:
        return "Hali tasdiqlangan ishchi yo'q."
    came, late, absent, waiting, on_leave = [], [], [], [], []
    expected = 0
    for w in people:
        sched = schedule(w, settings)
        branch = f" [{names[w['branch_id']]}]" if w["branch_id"] in names else ""
        night = db.latest_open_shift(w["tg_id"])
        if night and night["date"] != day.isoformat():
            came.append(f"• {w['name']}{branch}: tungi smena, {night['in_at'][8:10]}.{night['in_at'][5:7]} "
                        f"{hm(night['in_at'])} dan beri ishda")
            continue
        kind = cal.kind(w, day)
        s = db.get_shift(w["tg_id"], day)
        if kind == "work":
            expected += 1
        if s:
            info = shift_info(s, settings, w, names)
            tail = f" → {hm(s['out_at'])}" if s["out_at"] else " (ishda)"
            mark = f" ⏰ +{info.late} daq" if info.late else ""
            came.append(f"• {w['name']}{branch}: {hm(s['in_at'])}{tail}{mark}")
            if info.late:
                late.append(w["name"])
        elif kind.startswith("leave:"):
            on_leave.append(f"{w['name']} ({LEAVE_KINDS[kind[6:]].lower()})")
        elif kind == "work":
            start = sched.window(day)[0]
            if minutes(start, current) > int(settings["grace"]):
                absent.append(w["name"])
            else:
                waiting.append(f"{w['name']} ({sched.start})")
    head = f"📋 <b>Bugun, {day:%d.%m.%Y} ({WEEKDAYS[day.weekday()]})</b>"
    if day.isoformat() in cal.holidays:
        head += f"  · 🎉 {cal.holidays[day.isoformat()]}"
    lines = [head, f"Keldi: <b>{len(came)}</b>, ish kuni bo'lganlar: {expected}", ""]
    lines += came or ["Hali hech kim kelmagan."]
    if late:
        lines += ["", f"⏰ Kechikkanlar ({len(late)}): " + ", ".join(late)]
    if absent:
        lines += ["", f"❌ Kelmaganlar ({len(absent)}): " + ", ".join(absent)]
    if waiting:
        lines += ["", "🕓 Smenasi hali boshlanmagan: " + ", ".join(waiting)]
    if on_leave:
        lines += ["", "🏖 Ta'til/ruxsatda: " + ", ".join(on_leave)]
    return "\n".join(lines)


# ---------- Excel ----------

HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
ABSENT_FILL = PatternFill("solid", fgColor="FCE4E4")
LEAVE_FILL = PatternFill("solid", fgColor="E2EFDA")


def _sheet(ws, headers: list[str], rows: list[list], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def month_excel(stats: list[WorkerMonth], start: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Umumiy"
    rows = []
    for wm in sorted(stats, key=lambda x: x.worker["name"]):
        w = wm.worker
        rows.append([w["name"], w["position"] or "", w["phone"], wm.branch, wm.sched_label, wm.workdays, wm.days, len(wm.absent_days),
                     wm.paid_leave, wm.unpaid_leave, wm.late_count, wm.late_min, wm.early_min,
                     round(wm.worked / 60, 2), w["rate_type"] or "", w["rate"] or "",
                     round(wm.pay) if wm.pay is not None else ""])
    _sheet(ws, ["Ishchi", "Lavozim", "Telefon", "Filial", "Grafik", "Ish kunlari", "Kelgan kunlar", "Kelmagan kunlar",
                "Ta'til/kasal (haq to'lanadi)", "Ruxsat (o'z hisobidan)", "Kechikishlar soni", "Kechikish (daq)",
                "Erta ketish (daq)", "Jami soat", "Maosh turi", "Stavka", "Hisoblangan maosh"],
           rows, [24, 16, 16, 14, 16, 9, 9, 9, 12, 12, 10, 10, 10, 9, 10, 12, 15])
    if rows:
        total_row = len(rows) + 2
        ws.cell(row=total_row, column=1, value="JAMI").font = Font(bold=True)
        for col in (7, 8, 9, 10, 11, 12, 13, 14, 17):
            letter = get_column_letter(col)
            ws.cell(row=total_row, column=col, value=f"=SUM({letter}2:{letter}{total_row - 1})").font = Font(bold=True)

    detail = wb.create_sheet("Batafsil")
    drows = []
    for wm in stats:
        name = wm.worker["name"]
        for s in wm.shifts:
            note = "avto yopilgan" if s.auto else s.note
            drows.append([s.day, name, s.branch, hm(s.in_at), hm(s.out_at) if s.out_at else "ishda",
                          round(s.worked / 60, 2), s.late, s.early, note])
        for d in wm.absent_days:
            drows.append([d, name, "", "", "", 0, "", "", "kelmagan"])
        for kind, days in wm.leave_days.items():
            for d in days:
                drows.append([d, name, "", "", "", 0, "", "", LEAVE_KINDS[kind]])
    drows.sort(key=lambda r: (r[1], r[0]))
    drows = [[r[0].strftime("%d.%m.%Y"), WEEKDAYS[r[0].weekday()], *r[1:]] for r in drows]
    _sheet(detail, ["Sana", "Hafta kuni", "Ishchi", "Filial", "Keldi", "Ketdi", "Soat", "Kechikish (daq)",
                    "Erta ketish (daq)", "Izoh"], drows, [12, 12, 24, 14, 8, 8, 8, 11, 11, 22])
    for row in detail.iter_rows(min_row=2):
        note = row[9].value
        fill = ABSENT_FILL if note == "kelmagan" else LEAVE_FILL if note in LEAVE_KINDS.values() else None
        if fill:
            for cell in row:
                cell.fill = fill

    holidays = db.holidays_between(start, (start + timedelta(days=32)).replace(day=1))
    if holidays:
        hs = wb.create_sheet("Bayramlar")
        _sheet(hs, ["Sana", "Bayram"], [[date.fromisoformat(d).strftime("%d.%m.%Y"), n] for d, n in holidays.items()],
               [12, 30])

    wb.properties.title = f"{COMPANY}: {month_title(start)}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
