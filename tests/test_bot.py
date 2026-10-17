"""Runs every flow against a fake Telegram session. Prints ALL OK at the end."""
import asyncio
import os
import sys
from datetime import datetime, timedelta

HERE = os.path.dirname(os.path.abspath(__file__))
os.environ["DB_PATH"] = os.path.join(HERE, "test.db")
os.environ["BOT_TOKEN"] = "123:abc"
os.environ["ADMIN_IDS"] = "1"
if os.path.exists(os.environ["DB_PATH"]):
    os.remove(os.environ["DB_PATH"])
sys.path.insert(0, os.path.dirname(HERE))

from aiogram import Bot, Dispatcher
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.types import CallbackQuery, Chat, Contact, Location, Message, PhotoSize, Update, User, VideoNote

from app import admin, db, report, scheduler, utils, worker

sent = []  # (chat_id, text or document, reply_markup)


class FakeSession(BaseSession):
    async def make_request(self, bot, method, timeout=None):
        name = type(method).__name__
        if name == "SendVideoNote":
            sent.append((method.chat_id, f"[video] {method.video_note}", None))
        if name in ("SendMessage", "SendDocument", "EditMessageText", "SendPhoto"):
            body = getattr(method, "text", None) or getattr(method, "document", None) or f"[photo] {getattr(method, 'caption', '')}"
            sent.append((getattr(method, "chat_id", None), body, getattr(method, "reply_markup", None)))
        if name in ("SendMessage", "SendDocument", "SendPhoto", "SendVideoNote"):
            return Message(message_id=1, date=datetime.now(), chat=Chat(id=1, type="private"), text="x")
        return True

    async def stream_content(self, *a, **k):
        yield b""

    async def close(self):
        pass


async def main():
    db.init()
    # a database from the first release kept one workplace in settings: it becomes a branch
    db.set_setting("lat", "41.3")
    db.set_setting("lon", "69.2")
    db.init()
    assert [b["name"] for b in db.branches()] == ["Asosiy"] and db.get_settings()["lat"] == ""
    db.delete_branch(db.branches()[0]["id"])
    bot = Bot("123:abc", session=FakeSession(), default=DefaultBotProperties(parse_mode="HTML"))
    dp = Dispatcher()
    dp.include_router(admin.router)
    dp.include_router(worker.router)
    n = 0

    async def feed(uid, **kw):
        nonlocal n
        n += 1
        sent.clear()
        msg = Message(message_id=n, date=datetime.now(), chat=Chat(id=uid, type="private"),
                      from_user=User(id=uid, is_bot=False, first_name=f"U{uid}"), **kw)
        await dp.feed_update(bot, Update(update_id=n, message=msg))
        label = kw.get("text") or next(iter(kw))
        print(f">>> [{uid}] {label}\n" + "\n".join(f"  -> {c}: {t}" for c, t, _ in sent) + "\n")
        return [str(t) for _, t, _ in sent]

    async def tap(uid, data):
        nonlocal n
        n += 1
        sent.clear()
        cb = CallbackQuery(id=str(n), chat_instance="c", data=data, from_user=User(id=uid, is_bot=False, first_name="A"),
                           message=Message(message_id=9, date=datetime.now(), chat=Chat(id=uid, type="private"), text="ariza"))
        await dp.feed_update(bot, Update(update_id=n, callback_query=cb))
        print(f">>> [{uid}] TAP {data}\n" + "\n".join(f"  -> {c}: {t}" for c, t, _ in sent) + "\n")
        return [str(t) for _, t, _ in sent]

    ADMIN, ALI, VALI = 1, 101, 102
    await feed(ADMIN, text="/start")
    out = await feed(ALI, text="/start")
    assert "telefon" in out[0]
    # typing the number instead of tapping the button gets a hint, not silence
    out = await feed(ALI, text="+998901112233")
    assert "tugmasini" in out[0]
    # someone else's contact is rejected
    out = await feed(ALI, contact=Contact(phone_number="998900000000", first_name="X", user_id=999))
    assert "o'zingizning" in out[0]
    out = await feed(ALI, contact=Contact(phone_number="998901112233", first_name="Ali", last_name="Valiyev", user_id=ALI))
    assert any("Yangi ishchi" in t for t in out)
    await feed(VALI, contact=Contact(phone_number="+998907778899", first_name="Vali", user_id=VALI))
    out = await feed(ALI, text=worker.BTN_IN)
    assert "tasdiqlanmagan" in out[0]
    # workers cannot use admin commands
    out = await feed(ALI, text="/hisobot")
    assert not any("Oylik hisobot" in t for t in out)
    out = await tap(ADMIN, f"w:ok:{ALI}")
    assert any("lavozim" in t for t in out) and db.get_worker(ALI)["status"] == "pending"
    out = await feed(ADMIN, text="Sotuvchi")
    assert db.get_worker(ALI)["status"] == "active" and db.get_worker(ALI)["position"] == "Sotuvchi"
    assert any("Sotuvchi</b> lavozimida" in t for t in out)
    await tap(ADMIN, f"w:ok:{VALI}")
    await feed(ADMIN, text="-")
    assert db.get_worker(VALI)["status"] == "active" and db.get_worker(VALI)["position"] is None
    out = await feed(ADMIN, text=f"/lavozim {VALI} Omborchi")
    assert db.get_worker(VALI)["position"] == "Omborchi"
    # backdate both workers so last month's days count
    with db.db() as c:
        c.execute("UPDATE workers SET created_at = '2026-09-01 08:00'")

    # check in / out without location requirement
    out = await feed(ALI, text=worker.BTN_IN)
    assert "Xush kelibsiz" in out[0] and any("Ali Valiyev keldi" in t for t in out)
    out = await feed(ALI, text=worker.BTN_IN)
    assert "kelgansiz" in out[0]
    out = await feed(ALI, text=worker.BTN_OUT)
    assert "Ketgan vaqt" in out[0]
    out = await feed(ALI, text=worker.BTN_OUT)
    assert "ketgansiz" in out[0]
    out = await feed(VALI, text=worker.BTN_OUT)
    assert "Avval" in out[0]

    # location requirement
    await feed(ADMIN, text="/filial qosh Asosiy")
    await feed(ADMIN, location=Location(latitude=41.311081, longitude=69.240562))
    assert len(db.branches()) == 1
    await feed(VALI, text=worker.BTN_IN)
    out = await feed(VALI, location=Location(latitude=41.330000, longitude=69.240562))  # ~2 km away
    assert "uzoqdasiz" in out[0] and db.get_shift(VALI, utils.today()) is None
    await feed(VALI, text=worker.BTN_IN)
    out = await feed(VALI, location=Location(latitude=41.311500, longitude=69.240800))  # ~50 m
    assert "Xush kelibsiz" in out[0]
    BR1 = db.branches()[0]["id"]
    await feed(ADMIN, text=f"/filial {BR1} ochir")
    assert not db.branches()

    # admin views and settings
    out = await feed(ADMIN, text=admin.BTN_TODAY)
    assert "Keldi: <b>2</b>" in out[0]
    await feed(ADMIN, text="/ishchilar")
    await feed(ADMIN, text=f"/maosh {ALI} oylik 4m")
    await feed(ADMIN, text=f"/maosh {VALI} soatlik 25k")
    await feed(ADMIN, text="/sozlama ish 09:00 18:00")
    await feed(ADMIN, text="/sozlama dam shanba,yakshanba")
    out = await feed(ADMIN, text="/sozlama dam xyz")
    assert "Kun nomlari" in out[0]
    await feed(ADMIN, text="/sozlama dam yakshanba")

    # September history via corrections: Ali on time every workday except 2 absences, Vali late 15 min
    d = datetime(2026, 9, 1).date()
    while d.month == 9:
        if d.weekday() != 6:
            if d.day not in (10, 11):
                db.set_shift(ALI, d, f"{d} 08:55", f"{d} 18:05", "")
            db.set_shift(VALI, d, f"{d} 09:15", f"{d} 17:30", "")
        d += timedelta(days=1)
    out = await feed(ADMIN, text=f"/tuzat {ALI} 12.09.2026 09:30 18:00")
    assert "09:30" in out[0]
    out = await feed(ADMIN, text=f"/tuzat {ALI} 14.09.2026 ochir")
    assert "o'chirildi" in out[0]
    out = await feed(ADMIN, text="/hisobot 2026-09")
    assert "Sentabr 2026" in out[0]
    xlsx = sent[-1][1]
    assert xlsx.filename == "davomat_2026-09.xlsx"
    with open(os.path.join(HERE, "davomat_2026-09.xlsx"), "wb") as f:
        f.write(xlsx.data)
    stats = {s.worker["tg_id"]: s for s in report.month_stats(*utils.parse_month("2026-09"))}
    ali, vali = stats[ALI], stats[VALI]
    print("Ali:", ali.days, "days,", len(ali.absent_days), "absent,", ali.late_count, "late,", round(ali.pay))
    print("Vali:", vali.days, "days,", vali.late_count, "late,", vali.late_min, "late min,", vali.early_min, "early,",
          round(vali.worked / 60, 1), "h,", round(vali.pay))
    assert ali.month_workdays == 26 and ali.days == 23 and len(ali.absent_days) == 3 and ali.late_count == 1
    assert round(ali.pay) == round(4_000_000 * 23 / 26)
    assert vali.days == 26 and vali.late_count == 26 and vali.late_min == 26 * 15 and vali.early_min == 26 * 30
    assert round(vali.pay) == round(25_000 * 26 * 8.25)
    await feed(ADMIN, text=f"/ishchi {VALI} 2026-09")
    await feed(VALI, text=worker.BTN_MINE)

    # scheduler: remind, summary, auto close
    real_now = scheduler.now
    with db.db() as c:
        c.execute("UPDATE shifts SET out_at = NULL WHERE tg_id = ? AND date = ?", (VALI, utils.today().isoformat()))
    scheduler.now = lambda: datetime.combine(utils.today(), datetime.min.time()).replace(hour=18, minute=45)
    sent.clear()
    await scheduler.tick(bot)
    texts = [str(t) for _, t, _ in sent]
    print(">>> TICK 18:45\n" + "\n".join(texts) + "\n")
    assert any("Ketdim" in t for t in texts)
    sent.clear()
    await scheduler.tick(bot)
    assert not any("Ish vaqti" in str(t) for _, t, _ in sent)
    scheduler.now = lambda: datetime.combine(utils.today() + timedelta(days=1), datetime.min.time()).replace(minute=5)
    sent.clear()
    await scheduler.tick(bot)
    texts = [str(t) for _, t, _ in sent]
    print(">>> TICK next day 00:05\n" + "\n".join(texts) + "\n")
    assert db.get_shift(VALI, utils.today())["auto_closed"] == 1
    scheduler.now = real_now

    # ---------- stage 2 ----------
    # holiday: 15.09 is no longer a workday; Ali's absence that day disappears
    await feed(ADMIN, text="/bayram 15.09.2026 Test bayrami")
    with db.db() as c:
        c.execute("DELETE FROM shifts WHERE tg_id = ? AND date = '2026-09-15'", (ALI,))
    # leave: Ali asks for 10-11.09 as paid ta'til, admin approves
    out = await feed(ALI, text="/ruxsat 10.09.2026 11.09.2026 tatil Oilaviy")
    lv = db.leaves_overlapping(*utils.parse_month("2026-09"), status="pending")[0]
    await tap(ADMIN, f"l:ok:{lv['id']}")
    out = await tap(ADMIN, f"l:ok:{lv['id']}")
    stats = {s.worker["tg_id"]: s for s in report.month_stats(*utils.parse_month("2026-09"))}
    ali = stats[ALI]
    print("Ali after holiday+leave:", ali.days, "days,", ali.absent_days, ali.leave_days, ali.month_workdays, round(ali.pay))
    assert ali.month_workdays == 25 and ali.days == 22 and len(ali.absent_days) == 1 and ali.paid_leave == 2
    assert round(ali.pay) == 4_000_000 * 24 // 25
    await feed(ADMIN, text=f"/tatil {VALI} 05.09.2026 ruxsat")
    stats = {s.worker["tg_id"]: s for s in report.month_stats(*utils.parse_month("2026-09"))}
    assert stats[VALI].unpaid_leave == 1 and stats[VALI].days == 26  # came anyway: shift wins, leave day still listed
    await feed(ADMIN, text="/tatillar 2026-09")
    await feed(ADMIN, text="/bayramlar")
    await feed(ADMIN, text="/hisobot 2026-09")
    from openpyxl import load_workbook
    import io as _io
    wb = load_workbook(_io.BytesIO(sent[-1][1].data))
    if os.getenv("SAVE_SAMPLE"):
        with open(os.getenv("SAVE_SAMPLE"), "wb") as f:
            f.write(sent[-1][1].data)
    assert wb.sheetnames == ["Umumiy", "Batafsil", "Bayramlar"]
    notes = [r[9] for r in wb["Batafsil"].iter_rows(min_row=2, values_only=True)]
    assert notes.count("Ta'til") == 2 and notes.count("kelmagan") == 1

    # night shift schedule
    SOBIR = 103
    await feed(SOBIR, contact=Contact(phone_number="998905554433", first_name="Sobir", user_id=SOBIR))
    await tap(ADMIN, f"w:ok:{SOBIR}")
    await feed(ADMIN, text="Qorovul")
    out = await feed(ADMIN, text=f"/grafik {SOBIR} 20:00 08:00")
    assert "tungi" in out[0]
    base = utils.today()
    real_wnow = worker.now
    worker.now = lambda: datetime.combine(base, datetime.min.time()).replace(hour=19, minute=55)
    out = await feed(SOBIR, text=worker.BTN_IN)
    assert "19:55" in out[0] and "Kechikish" not in out[0]
    worker.now = lambda: datetime.combine(base + timedelta(days=1), datetime.min.time()).replace(hour=7, minute=50)
    out = await feed(SOBIR, text=worker.BTN_OUT)
    assert "11 soat 55 daq" in out[0] and "Erta ketish: 10 daq" in out[0], out
    # arriving at 00:40 after a missed start belongs to the previous evening's shift
    with db.db() as c:
        c.execute("DELETE FROM shifts WHERE tg_id = ?", (SOBIR,))
    worker.now = lambda: datetime.combine(base, datetime.min.time()).replace(hour=0, minute=40)
    out = await feed(SOBIR, text=worker.BTN_IN)
    sh = db.latest_open_shift(SOBIR)
    assert sh["date"] == (base - timedelta(days=1)).isoformat() and "Kechikish: 4 soat 40 daq" in out[0], out
    worker.now = real_wnow
    out = await feed(ADMIN, text=f"/tuzat {SOBIR} 01.09.2026 20:00 08:00")
    assert "ertasi kuni" in out[0]
    assert db.get_shift(SOBIR, datetime(2026, 9, 1).date())["out_at"] == "2026-09-02 08:00"

    # branches + selfie: Ali is bound to Yunusobod, so Chilonzor is rejected
    await feed(ADMIN, text="/filial qosh Chilonzor")
    await feed(ADMIN, location=Location(latitude=41.2756, longitude=69.2034))
    await feed(ADMIN, text="/filial qosh Yunusobod")
    await feed(ADMIN, location=Location(latitude=41.3650, longitude=69.2870))
    yunusobod = [b for b in db.branches() if b["name"] == "Yunusobod"][0]["id"]
    await feed(ADMIN, text=f"/bogla {ALI} {yunusobod}")
    await feed(ADMIN, text="/filiallar")
    await feed(ADMIN, text="/sozlama selfi on")
    await feed(ADMIN, text=f"/tuzat {ALI} {utils.today():%d.%m.%Y} ochir")
    await feed(ALI, text=worker.BTN_IN)
    out = await feed(ALI, location=Location(latitude=41.2757, longitude=69.2035))
    assert "Yunusobod" in out[0] and "uzoqdasiz" in out[0]
    await feed(ALI, text=worker.BTN_IN)
    out = await feed(ALI, location=Location(latitude=41.3651, longitude=69.2871))
    assert "selfi" in out[0]
    out = await feed(ALI, photo=[PhotoSize(file_id="selfie1", file_unique_id="u1", width=10, height=10)])
    assert "Xush kelibsiz" in out[0] and any(t.startswith("[photo]") and "Yunusobod" in t for t in out), out
    assert db.get_shift(ALI, utils.today())["in_photo"] == "selfie1"
    out = await feed(ADMIN, text=admin.BTN_TODAY)
    assert "[Yunusobod]" in out[0]

    # round video mode: a photo is refused, the circle goes to the admin with place, distance and map link
    await feed(ADMIN, text="/sozlama selfi video")
    await feed(ALI, text=worker.BTN_OUT)
    out = await feed(ALI, location=Location(latitude=41.3651, longitude=69.2871))
    assert "dumaloq video" in out[0]
    out = await feed(ALI, photo=[PhotoSize(file_id="selfie2", file_unique_id="u2", width=10, height=10)])
    assert "dumaloq video kerak" in out[0]
    out = await feed(ALI, video_note=VideoNote(file_id="circle1", file_unique_id="c1", length=240, duration=5))
    assert "Xayr" in out[0] and "[video] circle1" in out
    assert any("Yunusobod" in t and " m " in t and "maps.google.com" in t for t in out), out
    assert db.get_shift(ALI, utils.today())["out_photo"] == "v:circle1"

    # block
    await feed(ADMIN, text=f"/bloklash {VALI}")
    out = await feed(VALI, text=worker.BTN_IN)
    assert "yopilgan" in out[0]
    os.remove(os.path.join(HERE, "davomat_2026-09.xlsx"))
    os.remove(os.environ["DB_PATH"])
    print("ALL OK")


asyncio.run(main())
